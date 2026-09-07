from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from pathlib import Path


OUT = Path(r"C:\dev\closedshopping\docs\딜키_광고·상품경험_기획서.docx")

INK = "1A1714"
MUTED = "6B6258"
GOLD = "B8925A"
GOLD_DARK = "8C6D3F"
PALE = "F6F1E8"
PALE_GOLD = "EFE6D3"
LINE = "E6DDCD"
RED = "B34E4E"
GREEN = "3E6B53"
BLUE = "355C7D"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6", val="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_shading(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        borders.append(left)
        p_pr.append(borders)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    run._r.addnext(fld)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, GOLD_DARK, 18, 10),
        ("Heading 2", 13, GOLD_DARK, 14, 7),
        ("Heading 3", 12, INK, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("DEALKEY  /  AD & PRODUCT EXPERIENCE PLAN")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("딜키 내부 기획 문서  ·  ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(fp, "PAGE")


def add_para(doc, text="", style="Normal", bold=False, color=None, italic=False, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=None, color=color or INK, bold=bold, italic=italic)
    return p


def add_rich_para(doc, runs, style="Normal", align=None, fill=None, border=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if fill:
        set_paragraph_shading(p, fill, border)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
    for item in runs:
        text, kwargs = item if isinstance(item, tuple) else (item, {})
        run = p.add_run(text)
        set_run_font(run, color=kwargs.get("color", INK), bold=kwargs.get("bold"), italic=kwargs.get("italic"), size=kwargs.get("size"))
    return p


def add_bullet(doc, text, level=0, compact=False):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
        p.paragraph_format.first_line_indent = Inches(-0.188)
    run = p.add_run(text)
    if compact:
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.1
        set_run_font(run, color=INK, size=10)
    else:
        set_run_font(run, color=INK)
    return p


def add_number(doc, text):
    """Add a manually numbered paragraph so each logical list restarts at 1."""
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, color=INK)
    return p


def add_callout(doc, label, title, body, fill=PALE, border=GOLD):
    p = doc.add_paragraph()
    set_paragraph_shading(p, fill, border)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(label.upper() + "  ")
    set_run_font(r, size=8.5, color=border, bold=True)
    r = p.add_run(title + "\n")
    set_run_font(r, size=11.5, color=INK, bold=True)
    r = p.add_run(body)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_table(doc, headers, rows, widths, header_fill=PALE_GOLD, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header = table.rows[0]
    repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, color=GOLD_DARK, bold=True)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, color=GOLD, bold=True)
    return p


def build():
    doc = Document()
    style_document(doc)

    # First-page masthead: memo_masthead pattern with a restrained Dealkey override.
    add_kicker(doc, "DEALKEY  /  INTERNAL PLANNING GUIDE")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("딜키 광고·상품 경험 기획서")
    set_run_font(r, size=26, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(15)
    r = p.add_run("인스타그램 콘텐츠부터 상품 상세, 구매자 후기, 딜키 코멘트까지")
    set_run_font(r, size=13.5, color=MUTED)

    add_table(
        doc,
        ["문서 목적", "대상", "상태", "기준일"],
        [["신규 기획의 공통 기준 만들기", "대표·운영·개발·콘텐츠 담당", "후속 기획용 기준안", "2026-09-03"]],
        [1900, 2500, 2600, 2360],
        header_fill="F2EEE7",
        font_size=9,
    )

    add_callout(
        doc,
        "핵심 원칙",
        "애니메이션보다 ‘왜 싼지’와 ‘실제로 싼지’를 먼저 증명한다.",
        "딜키는 누구나 들어오는 오픈몰이 아니라 초대코드로 연결된 폐쇄몰이다. 따라서 첫 화면과 광고의 임무는 화려함보다 신뢰 형성, 가격의 근거, 가입 후 얻는 이익을 짧고 명확하게 보여주는 것이다.",
        fill=PALE_GOLD,
        border=GOLD_DARK,
    )

    add_para(doc, "이 문서는 무엇을 만들지뿐 아니라, 어떤 표현을 써야 하는지, 어떤 기능은 지금 미루는지, 누가 다음 결정을 내려야 하는지를 한 번에 이어받기 위한 기준 문서다.", after=7)
    add_para(doc, "문서 표기법: ‘확정’은 대화에서 이미 합의된 요구, ‘제안’은 실행을 빠르게 하기 위한 설계안, ‘검증 필요’는 실제 공급·가격·법적 근거 확인 전에는 광고에 쓰면 안 되는 항목이다.", color=MUTED, italic=True, after=10)

    add_para(doc, "목차", style="Heading 1")
    for item in [
        "1. 브랜드 포지셔닝과 커뮤니케이션 원칙",
        "2. 인스타그램 광고·캡션 운영 가이드",
        "3. 상품별 콘텐츠 플레이북",
        "4. 고객몰 사용자 경험과 신뢰 장치",
        "5. 후기와 ‘딜키 코멘트’ 기능 설계",
        "6. 관리자 운영 도구와 데이터 구조",
        "7. 단계별 개발·마케팅 로드맵",
        "8. 측정 지표와 실험 설계",
        "9. 신규 기획을 이어가기 위한 결정 질문",
    ]:
        add_bullet(doc, item, compact=True)

    doc.add_page_break()

    add_para(doc, "1. 브랜드 포지셔닝과 커뮤니케이션 원칙", style="Heading 1")
    add_para(doc, "1.1 딜키가 고객에게 약속하는 것", style="Heading 2")
    add_para(doc, "딜키의 차별점은 ‘예쁜 상품을 많이 보여주는 곳’이 아니라, 신뢰할 수 있는 사람이 소개한 상품을 제한된 회원에게 더 나은 조건으로 연결하는 구조다. 광고와 화면은 아래 세 문장을 일관되게 증명해야 한다.")
    for text in [
        "초대받은 사람만 들어오는 특판 구조: 추천코드 → 가입 승인 → 회원가·주문 권한의 흐름이 명확해야 한다.",
        "가격이 내려가는 이유가 있다: B2B 잔여 물량, 기획전 종료, 공급사 직연결 등은 실제 근거가 확인된 경우에만 사용한다.",
        "구매 후에도 책임지는 운영: 배송 상태, 문의 창구, 구매자 후기, 공급자·딜키의 설명이 남는다.",
    ]:
        add_bullet(doc, text)

    add_para(doc, "1.2 말투와 표현의 기준", style="Heading 2")
    add_table(
        doc,
        ["원칙", "권장 표현", "피할 표현", "이유"],
        [
            ["사실 우선", "‘B2B 판매 종료 후 남은 수량을 확인해 판매합니다’", "‘무조건 최저가’", "비교 기준과 수량 근거가 있어야 함"],
            ["폐쇄몰 명확화", "‘초대코드가 있는 분만 회원가를 확인할 수 있어요’", "‘아무나 가입 가능’", "서비스 구조와 광고 기대를 일치"],
            ["가격 구분", "‘온라인 기준가’ / ‘회원가’", "‘온라인가’만 단독 노출", "비회원이 결제 가격으로 오해할 수 있음"],
            ["절제된 희소성", "‘현재 확인된 수량 소진 시 종료’", "근거 없는 ‘마감 임박’", "가짜 긴급성은 신뢰를 훼손"],
            ["실제 후기", "‘구매자 후기’와 작성 시점·상품명 표시", "가공된 후기·AI 후기", "후기는 신뢰 자산이므로 출처가 필요"],
        ],
        [1350, 2950, 2500, 2560],
        font_size=8.8,
    )

    add_callout(
        doc,
        "검증 필요",
        "‘B2B 판매 후 남은 물량’은 상품마다 근거를 확인한 뒤 사용",
        "공급사가 실제로 어떤 판매를 마쳤는지, 잔여 수량이 무엇인지, 소비자 판매가 가능한지 확인한다. 근거가 없으면 ‘공급사 직연결 특판’, ‘기간 한정 공급가’처럼 확인된 사실만 사용한다.",
        fill="FFF7EC",
        border=RED,
    )

    add_para(doc, "1.3 고객의 머릿속에 남겨야 할 한 문장", style="Heading 2")
    add_rich_para(doc, [("딜키는 ", {"bold": True}), ("초대코드로 연결된 회원에게, 공급사가 확인한 상품을 더 합리적인 조건으로 보여주는 폐쇄형 특판몰", {"bold": True, "color": GOLD_DARK}), ("이다.", {})], fill=PALE, border=GOLD)
    add_para(doc, "이 한 문장은 홈 배너, 릴스 자막, 가입 안내, 상품 상세의 딜키 코멘트에서 표현만 바꿔 반복한다. 매체마다 전혀 다른 슬로건을 만들면 신규 방문자가 구조를 이해하는 데 시간이 더 걸린다.")

    add_para(doc, "2. 인스타그램 광고·캡션 운영 가이드", style="Heading 1")
    add_para(doc, "2.1 영상 한 편의 역할을 먼저 정한다", style="Heading 2")
    add_table(
        doc,
        ["콘텐츠 유형", "고객 질문", "영상에서 보여줄 것", "CTA"],
        [
            ["브랜드 입장 영상", "여기는 왜 폐쇄몰인가?", "초대코드·회원가·상품 검수 흐름", "프로필 링크에서 코드 확인"],
            ["가격 증명 영상", "정말 싼가?", "온라인 기준가 취소선 → 회원가, 가격 근거", "회원가 확인하기"],
            ["상품 사용 영상", "이 상품이 내게 필요한가?", "크기·사용 장면·포장·구성", "상세 정보 확인"],
            ["공급자/딜키 스토리", "아무거나 가져온 건 아닌가?", "공급자 설명·검수 포인트·딜키 코멘트", "상품 설명 읽기"],
            ["재고/마감 영상", "언제까지 살 수 있나?", "실제 잔여 수량·판매 종료 조건", "남은 수량 확인"],
        ],
        [1900, 2200, 3100, 2160],
        font_size=9,
    )

    add_para(doc, "2.2 릴스 기본 편집 구조(15~25초)", style="Heading 2")
    for i, text in enumerate([
        "0~2초 - 훅: ‘회원가가 따로 있는 이유, 15초 안에 보여드릴게요.’처럼 고객의 질문을 먼저 말한다.",
        "2~6초 - 맥락: ‘딜키는 초대코드가 있어야 들어오는 폐쇄몰입니다.’를 한 문장으로 설명한다.",
        "6~16초 - 증명: 상품 실물, 포장, 구성, 온라인 기준가와 회원가, 가격이 내려가는 근거를 순서대로 보여준다.",
        "16~22초 - 다음 행동: ‘프로필 링크에서 초대코드 승인 → 회원가입 → 회원가 확인’처럼 한 가지 행동만 요청한다.",
        "마지막 1~2초 - 기억 장치: DEALKEY 로고와 ‘좋은 상품을, 더 특별한 조건으로’ 같은 고정 문구를 남긴다.",
    ], 1):
        add_number(doc, f"{i}. {text}")

    add_para(doc, "2.3 캡션 공식", style="Heading 2")
    add_callout(
        doc,
        "캡션 템플릿",
        "질문 → 사실 → 가격 구분 → 가입 흐름 → 주의사항",
        "[질문형 첫 문장]\n[이 상품을 고른 이유와 확인된 사실 1~2문장]\n온라인 기준가: [금액] / 회원가: 가입 승인 후 확인\n초대코드가 있다면 프로필 링크에서 가입 승인을 시작하세요.\n※ 가격·재고·배송 조건은 게시 시점 기준이며 실제 화면을 확인해 주세요.",
        fill=PALE_GOLD,
        border=GOLD_DARK,
    )
    add_para(doc, "캡션은 영상 자막을 반복하는 공간이 아니다. 영상에서 보지 못한 가격 조건, 가입 순서, 배송·재고의 기준을 보충하는 공간으로 쓴다.")

    add_para(doc, "2.4 바로 사용할 수 있는 캡션 초안", style="Heading 2")
    captions = [
        ("브랜드/가입", "누구나 들어오는 쇼핑몰이 아닙니다.\n초대코드로 연결된 회원에게만 딜키의 회원가와 주문 기능이 열립니다.\n프로필 링크에서 코드 승인 → 회원가입 순서로 시작해 보세요.\n#딜키 #초대코드 #회원전용특판"),
        ("육포 선물세트", "명절 선물, 가격을 낮추는 이유가 있어야 하니까요.\n[공급사에서 확인한 사실을 한 문장으로 입력]\n온라인 기준가 [금액]은 비교용으로 표시하고, 회원가는 가입 승인 후 확인할 수 있습니다.\n구성·중량·배송 조건은 상세페이지에서 확인해 주세요."),
        ("생활용품", "매일 쓰는 물건일수록 구성과 사용 장면을 먼저 보여드릴게요.\n[원산지·용량·구성 등 확인된 정보]\n비회원에게는 온라인 기준가, 승인 회원에게는 회원가가 표시됩니다.\n초대코드가 있다면 프로필 링크에서 시작하세요."),
        ("가격 증명", "‘회원가’라고 쓰여 있으면 정말 다른 가격일까요?\n온라인 기준가와 회원가를 같은 상품·같은 구성으로 비교했습니다.\n가격은 상품 상세에서, 가입 승인은 프로필 링크에서 확인해 주세요."),
        ("공급자 스토리", "아무 상품이나 올리지 않습니다.\n[공급자가 직접 확인한 상품의 강점과 판매 조건]\n딜키 코멘트에서 이 상품을 고른 이유를 더 자세히 읽어보세요."),
    ]
    add_table(doc, ["용도", "캡션 초안"], captions, [1600, 7760], header_fill="F2EEE7", font_size=9)

    add_para(doc, "2.5 해시태그·댓글·링크 운영", style="Heading 2")
    for text in [
        "해시태그는 5~8개로 제한한다. 브랜드(#딜키), 구조(#초대코드 #회원전용), 상품군(#육포선물세트 등), 상황(#명절선물)으로 묶는다.",
        "첫 댓글에는 가입 순서와 주의사항을 고정한다: ‘초대코드 승인 → 회원가입 → 로그인 후 회원가 확인’.",
        "릴스마다 추천코드/캠페인 식별자를 분리한다. 링크에는 utm_source=instagram, utm_medium=reels, utm_campaign=[캠페인명]을 붙이고, 회원가입 시 추천 코드 귀속이 유지되는지 확인한다.",
        "댓글 질문은 가격·배송·가입·상품 정보로 분류해 답변 템플릿을 만든다. 확인되지 않은 재고·최저가·배송일은 즉답하지 않는다.",
    ]:
        add_bullet(doc, text)

    doc.add_page_break()
    add_para(doc, "3. 상품별 콘텐츠 플레이북", style="Heading 1")
    add_para(doc, "상품이 늘어나도 게시물 품질이 흔들리지 않도록 ‘상품의 사실 → 고객 장면 → 가격 근거 → 다음 행동’의 4칸으로 기획한다. 아래는 현재 대화에 등장한 상품군을 기준으로 한 초안이며, 대괄호는 등록 전에 채워야 한다.")

    add_para(doc, "3.1 The소프트 수제 육포 선물세트(600g·480g·300g)", style="Heading 2")
    add_table(
        doc,
        ["항목", "기획안", "등록 전 검증"],
        [
            ["핵심 훅", "‘선물용인데, 왜 회원가가 따로 있을까요?’", "회원가·온라인 기준가가 같은 구성인지"],
            ["증명 장면", "박스 전체 → 개별 포장 → 중량 표기 → 쇼핑백 → 실제 크기 비교", "실제 촬영 이미지, 구성품 수량"],
            ["가격 문장", "‘온라인 기준가 [금액] / 회원가 가입 승인 후 확인’", "비교 기준, 할인율 근거, 기간"],
            ["딜키 코멘트", "‘B2B/명절 기획 종료 후 남은 물량’은 사실 확인 시에만 사용", "공급자 확인 문서 또는 담당자 메모"],
            ["CTA", "‘초대코드 승인 후 300g·480g·600g 구성 비교하기’", "재고·배송비·묶음 조건"],
        ],
        [1500, 4980, 2880],
        font_size=9,
    )

    add_para(doc, "3.2 생활용품·세정제 상품", style="Heading 2")
    add_table(
        doc,
        ["고객이 궁금한 것", "영상/상세에서 보여줄 것", "주의할 점"],
        [
            ["어디에 쓰나?", "사용 장소, 용량, 사용 전·후 장면", "효능을 과장하거나 의학적 표현을 쓰지 않기"],
            ["구성이 무엇인가?", "본품·리필·세트 수량을 펼쳐서 보여주기", "사진과 실제 구성 불일치 금지"],
            ["왜 특판인가?", "공급자 직연결·기획전 조건 등 확인된 근거", "‘최저가’ 단정 대신 비교 기준 공개"],
        ],
        [2500, 4300, 2560],
        font_size=9,
    )

    add_para(doc, "3.3 신규 상품 등록용 1페이지 브리프", style="Heading 2")
    add_callout(
        doc,
        "등록 전 필수",
        "상품을 올리기 전에 콘텐츠 한 편의 답을 미리 적는다.",
        "상품명 / 카테고리 / 한 줄 장점 / 실제 구성 / 원산지·제조 정보 / 온라인 기준가 / 회원가 / 배송비·묶음 조건 / 재고 / 판매 중지 조건 / 가격 근거 / 공급자 코멘트 / 광고에서 금지할 표현",
        fill="F4F7F2",
        border=GREEN,
    )
    add_para(doc, "이 브리프가 비어 있으면 상품은 등록할 수 있어도 광고·상세·후기 운영으로 이어지지 않는다. 관리자 상품 등록 화면의 입력 항목과 이 브리프의 항목을 점차 맞춘다.")

    add_para(doc, "4. 고객몰 사용자 경험과 신뢰 장치", style="Heading 1")
    add_para(doc, "4.1 권장 고객 흐름", style="Heading 2")
    for i, text in enumerate([
        "릴스/추천 링크 유입: 영상에서 말한 한 가지 약속이 홈 배너와 첫 상품 섹션에서 바로 이어진다.",
        "상품 탐색: 카테고리 칩 → 상품 카드 → 상세페이지. 카드에는 썸네일, 상품명, 온라인 기준가(빨간 취소선), ‘회원가입 후 회원가 확인’을 표시한다.",
        "가입 승인: 초대코드 검증이 먼저이고, 승인 성공 후 회원가입 폼으로 이동한다. 비밀번호 확인 입력을 포함한다.",
        "구매: 로그인·추천 귀속·회원가 확인 → 배송지 선택/주소 검색 → 보내는 사람(기본 딜키, 수정 가능) → 결제 → 주문·배송 조회.",
        "구매 후: 배송완료 이후 후기 작성 요청. 상품 상세에는 구매자 후기와 딜키 코멘트를 구분해 보여준다.",
    ], 1):
        add_number(doc, f"{i}. {text}")

    add_para(doc, "4.2 가격과 배송을 오해 없이 보이기", style="Heading 2")
    add_table(
        doc,
        ["화면", "비회원", "승인 회원", "공통"],
        [
            ["상품 카드", "온라인 기준가(빨간 취소선) + 회원가입 안내", "회원가 + 필요 시 기준가 비교", "가격 옆에 배송 조건을 숨기지 않기"],
            ["상세 상단", "회원가 잠금 안내, 온라인 기준가", "회원가, 옵션·재고", "회원가·온라인 기준가의 기준일"],
            ["장바구니/결제", "로그인·가입 승인 유도", "회원가·상품별 배송비·묶음 조건", "결제 직전 총액과 배송지"],
        ],
        [1500, 2700, 2700, 2460],
        font_size=8.8,
    )
    add_callout(doc, "제안", "비회원에게는 ‘온라인 기준가’를 결제 가격처럼 보이지 않게 한다.", "빨간색 취소선과 ‘회원가입 후 회원가 확인’을 함께 사용한다. 단, 회원가가 아직 정해지지 않은 상품은 숫자를 임의로 보여주지 말고 ‘기준가 준비 중’으로 표시한다.", fill="FFF7EC", border=RED)

    add_para(doc, "5. 후기와 ‘딜키 코멘트’ 기능 설계", style="Heading 1")
    add_para(doc, "5.1 기능을 둘로 나누는 이유", style="Heading 2")
    add_para(doc, "구매자 후기는 실제 사용 경험의 증거이고, 딜키 코멘트는 상품을 왜 소싱하고 어떤 조건으로 소개하는지 설명하는 편집 영역이다. 두 내용을 한 칸에 섞으면 광고성 문구와 구매자 경험을 구분하기 어렵다.")
    add_table(
        doc,
        ["구분", "작성자", "표시 위치", "신뢰 규칙"],
        [
            ["구매자 후기", "실제 주문 회원", "상품 상세 ‘구매자 후기’", "배송완료 주문의 상품만 작성 가능, 주문상품·작성일 표시"],
            ["딜키 코멘트", "딜키 운영자", "상세 접힘 영역 안 별도 카드", "작성자 ‘딜키’·작성일·검증된 사실과 의견 구분"],
            ["공급자 코멘트", "공급자/운영자 입력", "딜키 코멘트 안 출처 블록", "공급자 확인 전에는 게시하지 않음"],
        ],
        [1700, 2000, 2700, 2960],
        font_size=8.8,
    )

    add_para(doc, "5.2 구매자 후기 MVP", style="Heading 2")
    for text in [
        "작성 자격: 주문 상태가 delivered인 order_item이 있는 회원. 같은 상품에 대한 후기 1개를 기본으로 하고, 수정은 허용한다.",
        "필수 입력: 별점(1~5), 후기 본문. 선택 입력: 사진 1~5장, 한 줄 요약.",
        "공개 상태: draft → published → hidden. 관리자 숨김 사유를 내부에 남긴다.",
        "표시 정보: ‘구매 인증’ 배지, 옵션/중량, 작성일, 닉네임 일부 마스킹. 전화번호·주소·이메일은 절대 노출하지 않는다.",
        "후기가 없을 때: ‘아직 구매자 후기가 없습니다. 첫 구매 후 경험을 남겨 주세요.’를 표시한다. 가짜 후기나 예시 후기를 실제 후기처럼 만들지 않는다.",
    ]:
        add_bullet(doc, text)

    add_para(doc, "5.3 딜키 코멘트 UX", style="Heading 2")
    for text in [
        "상품 상세의 구매 버튼 아래 또는 상세 정보 상단에 ‘딜키가 이 상품을 고른 이유’ 접힘 카드를 둔다.",
        "카드가 열리면 딜키 캐릭터/로고, 짧은 코멘트, 사실 블록(원산지·구성·가격 조건), 공급자 확인 메모를 순서대로 보여준다.",
        "코멘트는 300~600자 내외로 유지하고, 감탄사보다 구체적인 노력·확인·조건을 쓴다.",
        "예시 톤: ‘아무 상품이나 올린 것이 아닙니다. [확인한 과정]을 거쳐, [판매 조건]이 맞는 상품만 이번 특판에 소개합니다.’",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "초안", "딜키 코멘트 예시", "‘추석 전용 B2B 판매가 끝난 뒤, 소비자에게도 합리적인 조건으로 소개할 수 있는 수량을 공급사와 다시 확인했습니다. 구성·중량·배송 조건을 먼저 공개하고, 회원가가 적용되는 구조를 투명하게 안내합니다.’", fill=PALE, border=GOLD_DARK)

    add_para(doc, "5.4 문의/챗봇은 후순위", style="Heading 2")
    add_para(doc, "문의사항과 챗봇은 운영시간, 배송, 환불, 상품별 예외 규칙을 충분히 정한 뒤 붙여야 한다. 지금은 상품 상세에 고정 문의 안내와 관리자 확인 상태만 제공하고, 챗봇은 후기·배송 데이터가 쌓인 이후 별도 기획으로 둔다.")

    doc.add_page_break()
    add_para(doc, "6. 관리자 운영 도구와 데이터 구조", style="Heading 1")
    add_para(doc, "6.1 상품 관리에 필요한 필드", style="Heading 2")
    add_table(
        doc,
        ["영역", "필드", "운영 목적", "권한/상태"],
        [
            ["가격", "온라인 기준가, 회원가", "비회원 비교 가격과 승인 회원 결제 가격 분리", "관리자만 수정, 변경 이력 기록"],
            ["진열", "카테고리, 카테고리별 순서, 노출", "홈 섹션과 상품 목록 제어", "카테고리별 순서, 판매중만 고객 홈 노출"],
            ["판매", "재고, 예약재고, 판매중/판매중지중", "품절 자동 중지와 고객 안내", "재고 0이면 판매중지, 수동 재개는 재고 확인 후"],
            ["근거", "원산지, 구성, 가격 근거, 공급자 메모", "광고·딜키 코멘트의 사실 기반", "게시 전 검증 필요 표시"],
            ["콘텐츠", "썸네일, 상세 이미지, 딜키 코멘트", "목록·상세·광고 소재 분리", "썸네일/상세 역할 분리"],
        ],
        [1300, 2450, 3300, 2460],
        font_size=8.7,
    )

    add_para(doc, "6.2 후기 관리자 화면", style="Heading 2")
    for text in [
        "기본 큐: 새 후기 / 신고·검토 필요 / 게시됨 / 숨김. 10만 회원을 고려해 페이지네이션과 상품·기간·별점·후기 상태 필터를 기본으로 둔다.",
        "상품 상세에서 ‘후기 없음’과 ‘후기 있음’을 구분해 소싱·마케팅 담당자가 다음 콘텐츠를 고를 수 있게 한다.",
        "딜키 코멘트는 상품별 한 개의 현재본과 수정 이력을 보관한다. 공급자 메모는 공개 코멘트와 별도 필드로 둔다.",
        "모든 공개/숨김/수정 작업에는 관리자, 시각, 사유를 audit log로 남긴다.",
    ]:
        add_bullet(doc, text)

    add_para(doc, "6.3 회원·추천·콘텐츠 측정", style="Heading 2")
    add_para(doc, "회원 관리 화면은 단순 전체 목록이 아니라 ‘가입월’, ‘추천코드/캠페인’, ‘결제 경험’, ‘후기 작성’, ‘이메일 수신 동의’, ‘최근 방문’으로 필터·태그할 수 있어야 한다. 다만 초기에는 저장된 태그와 집계 필터만 제공하고, 복잡한 CRM 자동화는 거래가 쌓인 뒤 확장한다.")
    add_table(
        doc,
        ["세그먼트", "활용 예", "필요 데이터"],
        [
            ["8월/9월 유입", "시즌 캠페인 성과 비교", "가입일, 최초 유입 캠페인"],
            ["결제 경험 있음/없음", "재구매·첫 구매 리마인드", "유효 결제 주문, 취소 제외 규칙"],
            ["후기 작성자", "후기 기반 콘텐츠·감사 메시지", "게시 후기 수, 마지막 작성일"],
            ["이메일 수신 동의", "땡처리·신규 상품 알림", "동의 시각, 철회 시각, 발송 이력"],
        ],
        [2300, 4200, 3010],
        font_size=9,
    )

    add_para(doc, "7. 단계별 개발·마케팅 로드맵", style="Heading 1")
    add_para(doc, "기능을 한 번에 크게 만들지 않고, 판매를 시작하는 데 필요한 신뢰 장치부터 붙인다. 아래 순서는 ‘기능 구현 → 실제 판매 → 데이터 축적 → 자동화’의 흐름이다.")
    add_table(
        doc,
        ["단계", "범위", "완료 기준", "우선순위"],
        [
            ["P0: 캠페인 시작", "릴스 캡션 템플릿, 추천코드/UTM, 상품별 1페이지 브리프, 가격 표현 통일", "영상 1편을 게시하고 유입·가입·상품상세를 추적", "즉시"],
            ["P1: 신뢰 가능한 상세", "딜키 코멘트, 구매자 후기 MVP, 관리자 후기 검수, 근거 필드", "첫 구매자가 배송완료 후 후기를 남기고 상세에 표시", "다음"],
            ["P2: 운영 안정화", "택배사 API, 배송 상태 자동화, 문의 접수 템플릿, 재고·품절 운영", "주문부터 송장·배송완료까지 관리자 수기 복사를 줄임", "판매량 발생 후"],
            ["P3: 성장 자동화", "회원 세그먼트, 이메일 캠페인, 협력사/셀러 포털, 고도화된 추천", "캠페인별 매출·재구매·후기 전환을 자동 리포트", "데이터 축적 후"],
        ],
        [1300, 3500, 3200, 1510],
        font_size=8.8,
    )

    add_para(doc, "7.1 P0 게시 체크리스트", style="Heading 2")
    for text in [
        "영상의 첫 2초에 고객 질문 또는 가격 이유가 보이는가?",
        "온라인 기준가와 회원가의 기준이 같은 상품·같은 구성인가?",
        "초대코드 승인 → 회원가입 → 로그인 흐름이 캡션과 실제 화면에서 같은가?",
        "상품 상세에 썸네일, 구성, 배송비, 재고 조건이 있는가?",
        "게시 링크에 캠페인 식별자가 붙고 추천 귀속이 유지되는가?",
        "확인되지 않은 최저가·마감·잔여 수량 표현이 없는가?",
    ]:
        add_bullet(doc, text)

    add_para(doc, "8. 측정 지표와 실험 설계", style="Heading 1")
    add_para(doc, "팔로워 수보다 ‘실제 구매로 이어지는 신뢰’를 본다. 지표는 퍼널 순서대로 수집하고, 한 번에 한 가지 가설만 바꾼다.")
    add_table(
        doc,
        ["단계", "핵심 이벤트", "해석 질문"],
        [
            ["도달", "3초 시청률, 완주율, 저장, 공유", "첫 문장이 고객 질문을 붙잡았는가?"],
            ["유입", "프로필 클릭, 링크 클릭, 추천코드 페이지 도달", "영상의 약속과 랜딩 화면이 연결되는가?"],
            ["가입", "코드 검증 성공, 가입 완료, 첫 로그인", "승인 흐름이 번거롭거나 이해하기 어려운가?"],
            ["탐색", "상품상세 조회, 가격 확인, 장바구니", "가격·상품 근거가 구매 판단에 도움이 되는가?"],
            ["구매/신뢰", "결제 성공, 배송완료, 후기 작성", "구매 후 경험이 다음 콘텐츠의 증거가 되는가?"],
        ],
        [1500, 4000, 4010],
        font_size=9,
    )
    add_callout(doc, "실험 예시", "같은 상품, 첫 문장만 바꿔 비교", "A안: ‘회원가가 따로 있는 이유를 보여드릴게요.’ / B안: ‘명절 선물세트 가격이 내려간 이유는 따로 있습니다.’ 나머지 영상·랜딩·추천코드는 동일하게 유지하고 코드 검증 성공률과 상품상세 진입률을 비교한다.", fill="F4F7F2", border=GREEN)

    add_para(doc, "9. 신규 기획을 이어가기 위한 결정 질문", style="Heading 1")
    add_para(doc, "다음 기획자는 아래 질문에 답을 채우면서 문서를 확장한다. 답이 없는 항목은 광고 문구나 기능으로 확정하지 않는다.")
    questions = [
        "상품이 싸지는 정확한 이유는 무엇인가? 공급사가 확인할 수 있는 문서·수량·기간이 있는가?",
        "온라인 기준가는 어떤 비교 기준이며, 회원가는 언제까지 유효한가?",
        "구매자 후기는 배송완료 후 며칠 안에 요청할 것인가? 사진·별점·수정 정책은?",
        "딜키 코멘트의 공개 책임자는 누구이며, 공급자 코멘트와 어떤 순서로 검수하는가?",
        "회원에게 보낼 이메일의 동의 문구·수신거부·발송 빈도는 무엇인가?",
        "택배사/3PL별 송장 전달 방식과 실패 시 수기 대체 절차는 무엇인가?",
        "챗봇이 답해야 할 운영시간·배송·환불 정책이 문서로 확정되어 있는가? 확정 전에는 FAQ로 대체할 수 있는가?",
        "셀러·협력사에게 필요한 정산·초대코드·권한을 일반 회원과 어떻게 구분할 것인가?",
    ]
    for i, q in enumerate(questions, 1):
        add_number(doc, f"{i}. {q}")

    add_para(doc, "부록 A. 게시물 제작 요청서", style="Heading 1")
    add_table(
        doc,
        ["항목", "작성 내용"],
        [
            ["캠페인/추천코드", "[예: 2026 추석 육포 / 코드명]"],
            ["상품", "[상품명·옵션·카테고리]"],
            ["한 줄 약속", "[왜 이 상품을 봐야 하는가]"],
            ["검증된 사실", "[원산지·구성·가격 근거·재고·배송]"],
            ["영상 훅", "[첫 2초 자막]"],
            ["CTA", "[가입 승인 / 상세 보기 / 후기 남기기 중 하나]"],
            ["금지 표현", "[최저가·마감·효능 등 확인되지 않은 주장]"],
            ["게시 후 확인", "[3초 시청률·링크 클릭·코드 검증·상세 진입·구매·후기]"],
        ],
        [2400, 6960],
        header_fill="F2EEE7",
        font_size=9.2,
    )

    add_para(doc, "부록 B. 한 줄 정리", style="Heading 1")
    add_rich_para(doc, [("지금 만들 것: ", {"bold": True, "color": GOLD_DARK}), ("가격과 공급 근거를 보여주는 상품 상세 + 딜키 코멘트 + 구매자 후기 MVP.", {})], fill=PALE, border=GOLD)
    add_rich_para(doc, [("나중에 만들 것: ", {"bold": True, "color": BLUE}), ("챗봇, 셀러 포털, 고도화 CRM, 자동 이메일·택배 연동은 실제 구매 데이터가 쌓인 뒤 우선순위를 재평가한다.", {})], fill="F3F6F9", border=BLUE)

    doc.core_properties.title = "딜키 광고·상품 경험 기획서"
    doc.core_properties.subject = "인스타그램 광고, 상품 상세, 후기, 딜키 코멘트 및 후속 개발 계획"
    doc.core_properties.author = "Dealkey"
    doc.core_properties.keywords = "Dealkey, 광고, 캡션, 상품 경험, 후기, 딜키 코멘트, 폐쇄몰"
    doc.core_properties.comments = "Internal planning guide"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
